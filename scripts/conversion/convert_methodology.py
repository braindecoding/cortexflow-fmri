#!/usr/bin/env python3
"""
Methodology Document Converter
Converts METODOLOGI.md to PDF and DOCX formats with proper formatting
"""

import os
import sys
from pathlib import Path
import subprocess

def check_dependencies():
    """Check if required dependencies are installed"""
    
    print("🔍 CHECKING CONVERSION DEPENDENCIES")
    print("=" * 50)
    
    dependencies = {
        'pandoc': 'pandoc --version',
        'python-docx': 'python -c "import docx; print(docx.__version__)"',
        'markdown': 'python -c "import markdown; print(markdown.__version__)"',
        'weasyprint': 'python -c "import weasyprint; print(weasyprint.__version__)"'
    }
    
    missing = []
    available = []
    
    for dep, cmd in dependencies.items():
        try:
            result = subprocess.run(cmd, shell=True, capture_output=True, text=True)
            if result.returncode == 0:
                available.append(dep)
                print(f"✅ {dep}: Available")
            else:
                missing.append(dep)
                print(f"❌ {dep}: Not available")
        except Exception as e:
            missing.append(dep)
            print(f"❌ {dep}: Error - {e}")
    
    return available, missing

def install_dependencies():
    """Install missing dependencies"""
    
    print("\n📦 INSTALLING CONVERSION DEPENDENCIES")
    print("=" * 50)
    
    packages = [
        'python-docx',
        'markdown',
        'weasyprint',
        'pypandoc'
    ]
    
    for package in packages:
        try:
            print(f"Installing {package}...")
            subprocess.run([sys.executable, '-m', 'pip', 'install', package], 
                         check=True, capture_output=True)
            print(f"✅ {package} installed successfully")
        except subprocess.CalledProcessError as e:
            print(f"❌ Failed to install {package}: {e}")
    
    # Try to install pandoc via conda if available
    try:
        subprocess.run(['conda', 'install', '-c', 'conda-forge', 'pandoc', '-y'], 
                     check=True, capture_output=True)
        print("✅ pandoc installed via conda")
    except:
        print("❌ pandoc installation failed - please install manually")
        print("   Download from: https://pandoc.org/installing.html")

def convert_to_docx():
    """Convert methodology to DOCX format"""
    
    print("\n📄 CONVERTING TO DOCX FORMAT")
    print("=" * 40)
    
    try:
        from docx import Document
        from docx.shared import Inches, Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.style import WD_STYLE_TYPE
        import markdown
        import re
        
        # Read the markdown file
        with open('METODOLOGI.md', 'r', encoding='utf-8') as f:
            content = f.read()
        
        # Create new document
        doc = Document()
        
        # Set document margins
        sections = doc.sections
        for section in sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)
        
        # Add title
        title = doc.add_heading('METODOLOGI PENELITIAN CORTEXFLOW', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Add subtitle
        subtitle = doc.add_paragraph('Kerangka Kerja Dekoding Neural dengan Validasi Silang yang Ditingkatkan')
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle_format = subtitle.runs[0].font
        subtitle_format.size = Pt(14)
        subtitle_format.italic = True
        
        # Add date
        date_para = doc.add_paragraph(f'Tanggal: {datetime.now().strftime("%d %B %Y")}')
        date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Add page break
        doc.add_page_break()
        
        # Process content line by line
        lines = content.split('\n')
        current_table = []
        in_code_block = False
        code_content = []
        
        for line in lines:
            line = line.strip()
            
            if not line:
                if not in_code_block:
                    doc.add_paragraph()
                continue
            
            # Handle code blocks
            if line.startswith('```'):
                if in_code_block:
                    # End code block
                    code_para = doc.add_paragraph('\n'.join(code_content))
                    code_para.style = 'Intense Quote'
                    code_content = []
                    in_code_block = False
                else:
                    # Start code block
                    in_code_block = True
                continue
            
            if in_code_block:
                code_content.append(line)
                continue
            
            # Handle headers
            if line.startswith('#'):
                level = len(line) - len(line.lstrip('#'))
                header_text = line.lstrip('#').strip()
                if level <= 3:
                    doc.add_heading(header_text, level)
                else:
                    para = doc.add_paragraph(header_text)
                    para.style = 'Heading 4'
                continue
            
            # Handle tables
            if '|' in line and not line.startswith('!['):
                current_table.append(line)
                continue
            else:
                if current_table:
                    # Process accumulated table
                    add_table_to_doc(doc, current_table)
                    current_table = []
            
            # Handle images
            if line.startswith('!['):
                # Extract image info
                match = re.match(r'!\[(.*?)\]\((.*?)\)', line)
                if match:
                    alt_text, img_path = match.groups()
                    para = doc.add_paragraph(f'[Gambar: {alt_text}]')
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    para.runs[0].font.italic = True
                continue
            
            # Handle bullet points
            if line.startswith('- ') or line.startswith('* '):
                text = line[2:].strip()
                # Remove markdown formatting
                text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)  # Bold
                text = re.sub(r'\*(.*?)\*', r'\1', text)      # Italic
                doc.add_paragraph(text, style='List Bullet')
                continue
            
            # Handle numbered lists
            if re.match(r'^\d+\.', line):
                text = re.sub(r'^\d+\.\s*', '', line)
                text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)  # Bold
                text = re.sub(r'\*(.*?)\*', r'\1', text)      # Italic
                doc.add_paragraph(text, style='List Number')
                continue
            
            # Regular paragraph
            if line and not line.startswith('---'):
                # Remove markdown formatting
                text = re.sub(r'\*\*(.*?)\*\*', r'\1', line)  # Bold
                text = re.sub(r'\*(.*?)\*', r'\1', text)      # Italic
                doc.add_paragraph(text)
        
        # Process any remaining table
        if current_table:
            add_table_to_doc(doc, current_table)
        
        # Save document
        output_file = 'METODOLOGI_CortexFlow.docx'
        doc.save(output_file)
        print(f"✅ DOCX saved: {output_file}")
        return True
        
    except ImportError as e:
        print(f"❌ Missing dependency: {e}")
        return False
    except Exception as e:
        print(f"❌ DOCX conversion failed: {e}")
        return False

def add_table_to_doc(doc, table_lines):
    """Add table to document"""
    
    if len(table_lines) < 2:
        return
    
    # Parse table
    rows = []
    for line in table_lines:
        if '|' in line and not line.strip().startswith('|---'):
            cells = [cell.strip() for cell in line.split('|')[1:-1]]
            if cells:
                rows.append(cells)
    
    if not rows:
        return
    
    # Create table
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = 'Table Grid'
    
    # Fill table
    for i, row_data in enumerate(rows):
        row = table.rows[i]
        for j, cell_data in enumerate(row_data):
            if j < len(row.cells):
                # Remove markdown formatting
                cell_text = re.sub(r'\*\*(.*?)\*\*', r'\1', cell_data)
                row.cells[j].text = cell_text
                
                # Make header row bold
                if i == 0:
                    row.cells[j].paragraphs[0].runs[0].font.bold = True

def convert_to_pdf_pandoc():
    """Convert to PDF using pandoc"""
    
    print("\n📄 CONVERTING TO PDF (PANDOC)")
    print("=" * 40)
    
    try:
        cmd = [
            'pandoc',
            'METODOLOGI.md',
            '-o', 'METODOLOGI_CortexFlow.pdf',
            '--pdf-engine=xelatex',
            '--variable', 'geometry:margin=1in',
            '--variable', 'fontsize=11pt',
            '--variable', 'documentclass=article',
            '--variable', 'lang=id-ID',
            '--toc',
            '--number-sections'
        ]
        
        result = subprocess.run(cmd, capture_output=True, text=True)
        
        if result.returncode == 0:
            print("✅ PDF created successfully: METODOLOGI_CortexFlow.pdf")
            return True
        else:
            print(f"❌ Pandoc failed: {result.stderr}")
            return False
            
    except FileNotFoundError:
        print("❌ Pandoc not found")
        return False
    except Exception as e:
        print(f"❌ PDF conversion failed: {e}")
        return False

def convert_to_pdf_weasyprint():
    """Convert to PDF using weasyprint"""
    
    print("\n📄 CONVERTING TO PDF (WEASYPRINT)")
    print("=" * 40)
    
    try:
        import markdown
        import weasyprint
        from datetime import datetime
        
        # Read markdown
        with open('METODOLOGI.md', 'r', encoding='utf-8') as f:
            md_content = f.read()
        
        # Convert to HTML
        md = markdown.Markdown(extensions=['tables', 'toc', 'codehilite'])
        html_content = md.convert(md_content)
        
        # Create full HTML document
        html_doc = f"""
        <!DOCTYPE html>
        <html lang="id">
        <head>
            <meta charset="UTF-8">
            <title>Metodologi Penelitian CortexFlow</title>
            <style>
                body {{
                    font-family: 'Times New Roman', serif;
                    font-size: 11pt;
                    line-height: 1.6;
                    margin: 1in;
                    color: #333;
                }}
                h1, h2, h3, h4, h5, h6 {{
                    color: #2c3e50;
                    margin-top: 1.5em;
                    margin-bottom: 0.5em;
                }}
                h1 {{ font-size: 18pt; text-align: center; }}
                h2 {{ font-size: 16pt; }}
                h3 {{ font-size: 14pt; }}
                table {{
                    border-collapse: collapse;
                    width: 100%;
                    margin: 1em 0;
                }}
                th, td {{
                    border: 1px solid #ddd;
                    padding: 8px;
                    text-align: left;
                }}
                th {{
                    background-color: #f2f2f2;
                    font-weight: bold;
                }}
                code {{
                    background-color: #f4f4f4;
                    padding: 2px 4px;
                    border-radius: 3px;
                    font-family: 'Courier New', monospace;
                }}
                pre {{
                    background-color: #f4f4f4;
                    padding: 1em;
                    border-radius: 5px;
                    overflow-x: auto;
                }}
                .title-page {{
                    text-align: center;
                    margin-top: 2in;
                }}
                @page {{
                    margin: 1in;
                    @bottom-right {{
                        content: counter(page);
                    }}
                }}
            </style>
        </head>
        <body>
            <div class="title-page">
                <h1>METODOLOGI PENELITIAN CORTEXFLOW</h1>
                <h2>Kerangka Kerja Dekoding Neural dengan Validasi Silang yang Ditingkatkan</h2>
                <p><strong>Tanggal: {datetime.now().strftime("%d %B %Y")}</strong></p>
            </div>
            <div style="page-break-before: always;"></div>
            {html_content}
        </body>
        </html>
        """
        
        # Convert to PDF
        weasyprint.HTML(string=html_doc).write_pdf('METODOLOGI_CortexFlow.pdf')
        print("✅ PDF created successfully: METODOLOGI_CortexFlow.pdf")
        return True
        
    except ImportError as e:
        print(f"❌ Missing dependency: {e}")
        return False
    except Exception as e:
        print(f"❌ PDF conversion failed: {e}")
        return False

def main():
    """Main conversion function"""
    
    print("📄 METODOLOGI DOCUMENT CONVERTER")
    print("=" * 60)
    print("Converting METODOLOGI.md to PDF and DOCX formats")
    print()
    
    # Check if source file exists
    if not Path('METODOLOGI.md').exists():
        print("❌ METODOLOGI.md not found!")
        return
    
    # Check dependencies
    available, missing = check_dependencies()
    
    if missing:
        print(f"\n⚠️  Missing dependencies: {', '.join(missing)}")
        install_deps = input("Install missing dependencies? (y/n): ").lower().strip()
        if install_deps == 'y':
            install_dependencies()
    
    print("\n🔄 STARTING CONVERSIONS")
    print("=" * 40)
    
    # Convert to DOCX
    docx_success = convert_to_docx()
    
    # Convert to PDF (try multiple methods)
    pdf_success = False
    
    if 'pandoc' in available:
        pdf_success = convert_to_pdf_pandoc()
    
    if not pdf_success and 'weasyprint' in available:
        pdf_success = convert_to_pdf_weasyprint()
    
    # Summary
    print("\n📊 CONVERSION SUMMARY")
    print("=" * 40)
    print(f"DOCX: {'✅ Success' if docx_success else '❌ Failed'}")
    print(f"PDF:  {'✅ Success' if pdf_success else '❌ Failed'}")
    
    if docx_success or pdf_success:
        print("\n🎉 CONVERSION COMPLETED!")
        print("Output files:")
        if docx_success:
            print("  📄 METODOLOGI_CortexFlow.docx")
        if pdf_success:
            print("  📄 METODOLOGI_CortexFlow.pdf")
    else:
        print("\n❌ CONVERSION FAILED")
        print("Please install required dependencies and try again")

if __name__ == "__main__":
    from datetime import datetime
    import re
    main()
